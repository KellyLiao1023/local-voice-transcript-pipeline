import os
import requests
import json
from docx import Document
from path_config import load_config

# --- 1. 路徑與金鑰設定（由 config.json 提供，見 path_config.py） ---
CFG = load_config()
NOTE_DIR       = CFG["dirs"]["notes"]
WORD_DIR       = CFG["dirs"]["reports"]
LM_API_TOKEN   = CFG["lm_api_token"]
LM_STUDIO_URL  = CFG["lm_studio_url"]

# --- 🧪 實驗參數：指定目標與模型 ---
TARGET_FILE_NAME = "標準錄音 108.txt"           # 👈 在這裡輸入你想單獨跑的檔名（需位於 Study_Notes 資料夾）
# CURRENT_MODEL    = "deepseek-r1-0528-qwen3-8b"      # 想測試的模型 ID
# MODEL_SUFFIX     = "deepseek_8b"          # 檔名後綴

CURRENT_MODEL    = "gpt-oss-20b"      # 想測試的模型 ID
MODEL_SUFFIX     = "gpt-oss-20b"          # 檔名後綴


# --- 2. 核心功能：呼叫 AI 並生成 Word (維持你要求的格式) ---
def call_ai_and_make_word(txt_path, file_base_name, model_id, suffix):
    with open(txt_path, "r", encoding="utf-8") as f:
        content = f.read()

    # prompt = f"""
    # 請將以下逐字稿整理成結構化的 Word 檔案：
    # 其內容盡量避免表格，重點可用條列的方式
    # 內容要包含以下兩點:
    # 標題：聊天主題
    # 對話大綱以及重點
    # 最後再用一段文字總結描述這段對話的目的例如屬於開會或是閒聊或是教學現場等
    # 並且用繁體中文回答
    prompt = f"""
    請將以下逐字稿整理成結構化的 Word 檔案
    用繁體中文回答

    逐字稿內容：
    {content}
    """
    
    payload = {
        "model": model_id,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.5
    }

    headers = {
        "Authorization": f"Bearer {LM_API_TOKEN}",
        "Content-Type": "application/json"
    }
    
    print(f"🧠 正在精準轉譯樣本：[{file_base_name}] 使用模型：[{suffix}]...")
    try:
        response = requests.post(LM_STUDIO_URL, json=payload, headers=headers, timeout=120000)
        res_data = response.json()
        
        if 'choices' not in res_data:
            print(f"❌ 異常：伺服器回傳錯誤。內容：{res_data}")
            return

        ai_reply = res_data['choices'][0]['message']['content']
        output_word = os.path.join(WORD_DIR, f"{file_base_name}_{suffix}.docx")
        
        # 安全機制：覆蓋確認
        if os.path.exists(output_word):
            choice = input(f"⚠️ {file_base_name}_{suffix}.docx 已存在，要覆蓋嗎？(y/n): ")
            if choice.lower() != 'y': return

        doc = Document()
        doc.add_heading(f'對話紀錄：{file_base_name} ({suffix})', 0)
        
        # 解析 AI Markdown 並填入 Word
        for line in ai_reply.split('\n'):
            line = line.strip()
            if not line: continue
            if line.startswith('### '): doc.add_heading(line.replace('### ', ''), level=3)
            elif line.startswith('## '): doc.add_heading(line.replace('## ', ''), level=2)
            elif line.startswith('# '): doc.add_heading(line.replace('# ', ''), level=1)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line.replace('- ', '').replace('* ', ''), style='List Bullet')
            else:
                doc.add_paragraph(line.replace('**', ''))
        
        doc.save(output_word)
        print(f"✅ 精準報告已生成：{output_word}")

    except Exception as e:
        print(f"❌ AI 摘要過程失敗：{e}")

# ==========================================
# 模式 D：單一檔案重處理模式 (Targeted Processing)
# ==========================================
def mode_single_file_reprocess(target_name):
    # 自動處理有無附檔名的情況
    base_name = target_name.replace(".txt", "")
    target_path = os.path.join(NOTE_DIR, f"{base_name}.txt")

    if os.path.exists(target_path):
        os.makedirs(WORD_DIR, exist_ok=True)
        call_ai_and_make_word(target_path, base_name, CURRENT_MODEL, MODEL_SUFFIX)
    else:
        print(f"❌ 錯誤：在 {NOTE_DIR} 找不到檔案 '{target_name}'")
        print("請確認檔案是否存在於 Study_Notes 資料夾中。")

# ==========================================
# 🧪 執行入口
# ==========================================
if __name__ == "__main__":
    # 直接執行單一檔案重處理
    mode_single_file_reprocess(TARGET_FILE_NAME)